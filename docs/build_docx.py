#!/usr/bin/env python3
"""Builds the same document as a Word file.

Imports the body from _content.py, so the .docx and the .pdf cannot describe
different projects. Only the rendering differs.
"""
import os
import re

import docx
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

import _content

OUT = os.path.join(os.path.dirname(__file__), 'Nuno_Project_Explained.docx')

INK = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT = RGBColor(0x5B, 0x3F, 0xA8)
MUTED = RGBColor(0x5A, 0x5A, 0x72)

doc = Document()

# Page setup: A4 with the same margins as the PDF.
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
for attr in ('left_margin', 'right_margin'):
    setattr(sec, attr, Inches(0.79))
sec.top_margin = Inches(0.71)
sec.bottom_margin = Inches(0.79)

normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10)
normal.font.color.rgb = INK


def shade(cell, hex_colour):
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:fill'), hex_colour)
    cell._tc.get_or_add_tcPr().append(el)


def runs_from(text, para, base_size=10, mono=False):
    """Renders `code` and **bold** markers into real runs."""
    token = re.compile(r'(`[^`]+`|\*\*[^*]+\*\*)')
    for part in token.split(text):
        if not part:
            continue
        if part.startswith('`') and part.endswith('`'):
            r = para.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(base_size - 1)
        elif part.startswith('**') and part.endswith('**'):
            r = para.add_run(part[2:-2])
            r.bold = True
            r.font.size = Pt(base_size)
        else:
            r = para.add_run(part)
            r.font.size = Pt(base_size)
            if mono:
                r.font.name = 'Consolas'


# ── API the shared content calls ─────────────────────────────────

def h1(t):
    para = doc.add_paragraph()
    r = para.add_run(t)
    r.bold = True
    r.font.size = Pt(19)
    r.font.color.rgb = INK
    para.space_after = Pt(8)


def h2(t):
    para = doc.add_paragraph()
    para.space_before = Pt(15)
    para.space_after = Pt(5)
    r = para.add_run(t)
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = ACCENT


def h3(t):
    para = doc.add_paragraph()
    para.space_before = Pt(10)
    para.space_after = Pt(3)
    r = para.add_run(t)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = INK


def p(t):
    para = doc.add_paragraph()
    para.space_after = Pt(6)
    runs_from(t, para)


def small(t):
    para = doc.add_paragraph()
    para.space_after = Pt(5)
    runs_from(t, para, base_size=9)
    for r in para.runs:
        r.font.color.rgb = MUTED


def code(t):
    # One shaded single-cell table: Word keeps the block visually together
    # and preserves the monospaced alignment.
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = tbl.cell(0, 0)
    shade(cell, 'F4F4F8')
    cell.text = ''
    for i, line in enumerate(t.split('\n')):
        para = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        para.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1.0
        r = para.add_run(line)
        r.font.name = 'Consolas'
        r.font.size = Pt(8.5)
    doc.add_paragraph().space_after = Pt(0)


def gap(h=4):
    para = doc.add_paragraph()
    para.space_after = Pt(max(1, int(h / 2)))


def bullets(items):
    for it in items:
        para = doc.add_paragraph(style='List Bullet')
        para.space_after = Pt(3)
        runs_from(it, para)


def qa(q, a):
    para = doc.add_paragraph()
    para.space_before = Pt(7)
    para.space_after = Pt(2)
    r = para.add_run('Q. ' + q)
    r.bold = True
    r.font.size = Pt(10)
    ans = doc.add_paragraph()
    ans.space_after = Pt(6)
    ans.paragraph_format.left_indent = Inches(0.18)
    runs_from(a, ans)


def table(rows, widths, header=True):
    tbl = doc.add_table(rows=0, cols=len(rows[0]))
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, row in enumerate(rows):
        cells = tbl.add_row().cells
        for j, text in enumerate(row):
            cells[j].text = ''
            para = cells[j].paragraphs[0]
            para.space_after = Pt(0)
            runs_from(text, para, base_size=9)
            if header and i == 0:
                for r in para.runs:
                    r.bold = True
                shade(cells[j], 'ECE8F7')
    doc.add_paragraph().space_after = Pt(0)


class _Story(list):
    """The shared content appends a PageBreak and one raw Paragraph."""

    def append(self, item):
        if item is _PAGE_BREAK:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif isinstance(item, _Sub):
            para = doc.add_paragraph()
            para.space_after = Pt(6)
            for line in item.text.split('\n'):
                if para.runs:
                    para.add_run().add_break()
                r = para.add_run(line)
                r.font.size = Pt(9)
                r.font.color.rgb = MUTED


_PAGE_BREAK = object()


class _Sub:
    def __init__(self, text):
        self.text = text


def _paragraph(text, style):
    # The subtitle is the only raw-markup paragraph in the content.
    clean = (text.replace('<br/>', '\n')
                 .replace('&#183;', '\u00b7')
                 .replace('&#160;', ' ')
                 .replace('&#8212;', '\u2014'))
    return _Sub(clean)


_content.build({
    'h1': h1, 'h2': h2, 'h3': h3, 'p': p, 'small': small, 'code': code,
    'gap': gap, 'bullets': bullets, 'qa': qa, 'table': table,
    'story': _Story(), 'W': 1.0, 'mm': 1.0,
    'Paragraph': _paragraph, 'SMALL': None, 'PageBreak': lambda: _PAGE_BREAK,
})

doc.save(OUT)
print(f'wrote {OUT} ({os.path.getsize(OUT):,} bytes)')
